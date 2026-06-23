import csv
from datetime import date, datetime, timedelta

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone

from .models import Bus, Trip, TripPassenger

MAX_RECURRING_MONTHS = 3
MAX_RECURRING_DAYS = MAX_RECURRING_MONTHS * 31


def has_available_bus(trip_date, route):
    """
    Checks if there is at least one active bus in the fleet that is not 
    allocated to a trip during the specified trip_date and route time.
    """
    active_buses = Bus.objects.filter(status="ATIVO")
    if not active_buses.exists():
        return False

    # Define a generous date range to find potentially overlapping trips
    date_range = [
        trip_date - timedelta(days=1),
        trip_date,
        trip_date + timedelta(days=1),
    ]
    overlapping_trips = Trip.objects.filter(
        trip_date__in=date_range
        ).exclude(bus__isnull=True)
    
    tz = timezone.get_current_timezone()
    new_start = timezone.make_aware(
        datetime.combine(trip_date, route.departure_time), tz)
    new_end = timezone.make_aware(
        datetime.combine(trip_date, route.arrival_time), tz)

    if new_end <= new_start:
        new_end += timedelta(days=1)

    busy_bus_ids = set()
    for existing_trip in overlapping_trips:
        ex_start = timezone.make_aware(
            datetime.combine(
                existing_trip.trip_date, existing_trip.route.departure_time
                ), tz
        )
        ex_end = timezone.make_aware(
            datetime.combine(
                existing_trip.trip_date, existing_trip.route.arrival_time
                ), tz
        )
        
        if ex_end <= ex_start:
            ex_end += timedelta(days=1)

        # Check for time intersection
        if new_start < ex_end and new_end > ex_start:
            busy_bus_ids.add(existing_trip.bus_id)
            
    # If there is any active bus that is not in the set of busy buses, return True
    return active_buses.exclude(id__in=busy_bus_ids).exists()


@transaction.atomic
def create_recurring_trips(date_start, date_end, weekdays, route, status):
    trips = []
    current = date_start
    while current <= date_end:
        if current.weekday() in weekdays:
            trips.append(
                Trip(trip_date=current, route=route, status=status)
            )
        current += timedelta(days=1)
    return Trip.objects.bulk_create(trips)


def export_trip_passengers(trip, fmt="csv"):
    """Gera arquivo CSV ou XLSX com a lista de passageiros da viagem."""
    from ..reservations.models import Reservation

    reservations = Reservation.objects.filter(trip=trip).select_related(
        "student__user", "civil_servant__user", "guest_passenger"
    )
    trip_passengers = trip.trip_passengers.select_related("allowed_staff")

    fieldnames = ["ID", "Passageiro", "Tipo", "Documento", "Status", "Check-in"]
    rows = []

    for r in reservations:
        if r.student:
            name = r.student.user.full_name
            ptype = "ESTUDANTE"
            doc = r.student.student_id
        elif r.civil_servant:
            name = r.civil_servant.user.full_name
            ptype = "SERVIDOR"
            doc = r.civil_servant.civil_servant_id
        elif r.guest_passenger:
            name = r.guest_passenger.full_name
            ptype = "CONVIDADO"
            doc = r.guest_passenger.cpf
        else:
            continue

        rows.append({
            "ID": r.id,
            "Passageiro": name,
            "Tipo": ptype,
            "Documento": str(doc or "—"),
            "Status": r.status,
            "Check-in": "Sim" if r.check_in else "Não",
        })

    for tp in trip_passengers:
        if tp.passenger_type == TripPassenger.PassengerType.LOCAL_SERVER and tp.allowed_staff:
            rows.append({
                "ID": tp.id,
                "Passageiro": tp.allowed_staff.name,
                "Tipo": "SERVIDOR LOCAL",
                "Documento": str(tp.allowed_staff.registration_number or "—"),
                "Status": "CONFIRMADA",
                "Check-in": "Sim",
            })
        elif tp.passenger_type == TripPassenger.PassengerType.LOCAL_GUEST:
            rows.append({
                "ID": tp.id,
                "Passageiro": tp.full_name,
                "Tipo": "CONVIDADO LOCAL",
                "Documento": str(tp.cpf or "—"),
                "Status": "CONFIRMADA",
                "Check-in": "Sim",
            })

    rows.sort(key=lambda p: p["ID"], reverse=True)

    if fmt == "xlsx":
        return _build_xlsx_response(trip.id, rows, fieldnames)

    return _build_csv_response(trip.id, rows, fieldnames)


def _build_csv_response(trip_id, rows, fieldnames):
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="viagem_{trip_id}_passageiros.csv"'
    writer = csv.DictWriter(response, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
    return response


def _build_xlsx_response(trip_id, rows, fieldnames):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Passageiros"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    for col, name in enumerate(fieldnames, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for r_idx, row in enumerate(rows, 2):
        for c_idx, key in enumerate(fieldnames, 1):
            ws.cell(row=r_idx, column=c_idx, value=row[key])

    for col in range(1, len(fieldnames) + 1):
        max_len = max(
            len(str(ws.cell(row=r, column=col).value or ""))
            for r in range(1, len(rows) + 2)
        )
        ws.column_dimensions[get_column_letter(col)].width = max_len + 2

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="viagem_{trip_id}_passageiros.xlsx"'
    wb.save(response)
    return response
